"""Generate Activity_Weight_Tracker.xlsx (formulas, validation, charts).

The VBA macros live in TrackerMacros.bas and are imported into Excel by the user.
Run: python3 build_tracker.py
"""
import datetime as dt
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference, ScatterChart, Series
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).with_name("Activity_Weight_Tracker.xlsx")
LAST = 1000  # last data row on Activities / Weight

FONT = "Arial"
NAVY = "1F3864"
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")   # light yellow = cells you fill in
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
CARD_FILL = PatternFill("solid", fgColor="DEEAF6")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
DATE_FMT = "dd/mm/yyyy"

# Activity, MET value (Compendium of Physical Activities, Ainsworth et al. 2011)
ACTIVITIES = [
    ("Walking", 3.5), ("Brisk walking", 4.3), ("Running", 9.8), ("Cycling", 7.5),
    ("Swimming", 6.0), ("Hiking", 6.0), ("Strength training", 5.0), ("Yoga", 2.5),
    ("Football", 7.0), ("Tennis", 7.3), ("Dancing", 5.0), ("Other", 4.0),
]
ACT_FIRST = 10
ACT_LAST = ACT_FIRST + len(ACTIVITIES) - 1
ACT_RANGE = f"Settings!$A${ACT_FIRST}:$A${ACT_LAST}"
MET_RANGE = f"Settings!$B${ACT_FIRST}:$B${ACT_LAST}"


def f(bold=False, color="000000", size=10, italic=False):
    return Font(name=FONT, bold=bold, color=color, size=size, italic=italic)


def header(ws, row, titles, widths):
    for i, (t, w) in enumerate(zip(titles, widths), start=1):
        c = ws.cell(row=row, column=i, value=t)
        c.font = f(bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[row].height = 30


wb = Workbook()

# ------------------------------------------------------------------ Instructions
ins = wb.active
ins.title = "Instructions"
ins.column_dimensions["A"].width = 110
lines = [
    ("Activity & Weight Tracker", f(bold=True, size=16, color=NAVY)),
    ("", None),
    ("1. Set up once", f(bold=True, size=12)),
    ("•  Settings sheet: enter your height, target weight and weekly activity goal (yellow cells).", None),
    ("•  Import the macros: press Alt+F11, then File > Import File... and choose TrackerMacros.bas. Close the editor.", None),
    ("•  File > Save As > 'Excel Macro-Enabled Workbook (*.xlsm)'. Reopen it and click 'Enable Content' if asked.", None),
    ("•  Press Alt+F8, run 'SetupButtons': two buttons appear on the Dashboard.", None),
    ("", None),
    ("2. Every day", f(bold=True, size=12)),
    ("•  Dashboard > '+ Add activity': date, activity, duration, distance (optional), notes.", None),
    ("•  Dashboard > '+ Add weight': date and weight. If a weight already exists for that date, you can replace it.", None),
    ("•  The macros sort the rows by date automatically.", None),
    ("•  Without macros you can also type directly in the yellow columns of the Activities and Weight sheets.", None),
    ("", None),
    ("3. What is calculated", f(bold=True, size=12)),
    ("•  Calories (estimate) = MET of the activity × your weight (kg) × duration (hours). MET values are in Settings.", None),
    ("•  The weight used is your last weight recorded on or before the activity date.", None),
    ("•  Weight sheet: change vs previous measurement, BMI, 7-measurement moving average, distance to target.", None),
    ("•  Dashboard: current weight, progress, this week vs your goal, last 30 days, totals per activity, charts.", None),
    ("", None),
    ("Legend", f(bold=True, size=12)),
    ("Yellow cells = data you enter. White cells = formulas, do not edit.", None),
    ("The first row of Activities and Weight is an example: overwrite or delete it.", None),
]
for i, (text, font) in enumerate(lines, start=1):
    c = ins.cell(row=i, column=1, value=text)
    c.font = font or f()
    c.alignment = Alignment(wrap_text=True, vertical="top")
ins["A22"].fill = INPUT_FILL
ins.sheet_view.showGridLines = False

# ---------------------------------------------------------------------- Settings
st = wb.create_sheet("Settings")
st.column_dimensions["A"].width = 32
st.column_dimensions["B"].width = 14
st.column_dimensions["C"].width = 60
st["A1"] = "Settings"
st["A1"].font = f(bold=True, size=14, color=NAVY)
settings = [
    (3, "Height (m)", 1.75, "0.00", "Your height in metres - used for BMI. Example value: replace it."),
    (4, "Target weight (kg)", 72, "0.0", "Your goal weight."),
    (5, "Weekly activity goal (min)", 150, "0", "WHO recommends at least 150 min of moderate activity per week."),
    (6, "Default weight (kg)", 75, "0.0", "Used for calories only until you record a first weight."),
]
for r, label, val, fmt, note in settings:
    st.cell(row=r, column=1, value=label).font = f(bold=True)
    c = st.cell(row=r, column=2, value=val)
    c.font = f(color="0000FF")
    c.fill = INPUT_FILL
    c.border = BOX
    c.number_format = fmt
    st.cell(row=r, column=3, value=note).font = f(italic=True, color="595959")

header(st, ACT_FIRST - 1, ["Activity", "MET", "Source / note"], [32, 14, 60])
for i, (name, met) in enumerate(ACTIVITIES):
    r = ACT_FIRST + i
    for col, v in ((1, name), (2, met)):
        c = st.cell(row=r, column=col, value=v)
        c.font = f(color="0000FF")
        c.fill = INPUT_FILL
        c.border = BOX
    st.cell(row=r, column=2).number_format = "0.0"
st.cell(row=ACT_FIRST, column=3,
        value="MET values: Compendium of Physical Activities (Ainsworth et al., 2011). "
              "You can rename activities or change MET values.").font = f(italic=True, color="595959")
st.cell(row=ACT_FIRST, column=3).alignment = Alignment(wrap_text=True, vertical="top")
st.merge_cells(start_row=ACT_FIRST, start_column=3, end_row=ACT_FIRST + 2, end_column=3)

wb.defined_names["ActivityTable"] = DefinedName(
    "ActivityTable", attr_text=f"Settings!$A${ACT_FIRST}:$B${ACT_LAST}")
wb.defined_names["ActivityList"] = DefinedName("ActivityList", attr_text=ACT_RANGE)

# -------------------------------------------------------------------- Weight
wt = wb.create_sheet("Weight")
header(wt, 1, ["Date", "Weight (kg)", "Change vs previous (kg)", "BMI",
               "Moving average (7 entries)", "To target (kg)"],
       [13, 13, 16, 9, 16, 14])
wt.freeze_panes = "A2"
for r in range(2, LAST + 1):
    for col in (1, 2):
        c = wt.cell(row=r, column=col)
        c.fill = INPUT_FILL
        c.font = f(color="0000FF")
        c.border = BOX
    wt.cell(row=r, column=1).number_format = DATE_FMT
    wt.cell(row=r, column=2).number_format = "0.0"
    prev = (f'=IF(OR(B{r}="",B{r-1}=""),"",B{r}-B{r-1})' if r > 2 else '=""')
    first = max(2, r - 6)
    formulas = [
        (3, prev, '+0.0;-0.0;0.0'),
        (4, f'=IF(B{r}="","",B{r}/Settings!$B$3^2)', "0.0"),
        (5, f'=IF(B{r}="","",AVERAGE(B{first}:B{r}))', "0.0"),
        (6, f'=IF(B{r}="","",B{r}-Settings!$B$4)', '+0.0;-0.0;0.0'),
    ]
    for col, formula, fmt in formulas:
        c = wt.cell(row=r, column=col, value=formula)
        c.font = f()
        c.number_format = fmt
        c.border = BOX
wt["A2"] = dt.date(2026, 10, 1)
wt["B2"] = 75.0
wt["A2"].comment = Comment("Example row - overwrite it with your own first weight.", "Tracker")
wt.conditional_formatting.add("C3:C1000", CellIsRule(operator="lessThan", formula=["0"],
                                                     font=Font(name=FONT, color="008000")))
wt.conditional_formatting.add("C3:C1000", CellIsRule(operator="greaterThan", formula=["0"],
                                                     font=Font(name=FONT, color="C00000")))
dv_date = DataValidation(type="date", operator="between", formula1="DATE(2000,1,1)",
                         formula2="DATE(2100,12,31)", showErrorMessage=True,
                         error="Enter a valid date.")
dv_w = DataValidation(type="decimal", operator="between", formula1="20", formula2="400",
                      showErrorMessage=True, error="Weight must be between 20 and 400 kg.")
wt.add_data_validation(dv_date)
wt.add_data_validation(dv_w)
dv_date.add(f"A2:A{LAST}")
dv_w.add(f"B2:B{LAST}")

# ----------------------------------------------------------------- Activities
ac = wb.create_sheet("Activities")
header(ac, 1, ["Date", "Activity", "Duration (min)", "Distance (km)", "Notes",
               "Week", "Pace (min/km)", "Weight used (kg)", "MET", "Calories (kcal, est.)"],
       [13, 20, 11, 11, 30, 8, 11, 12, 7, 13])
ac.freeze_panes = "A2"
W = "Weight!$A$2:$A$1000"
WB = "Weight!$B$2:$B$1000"
for r in range(2, LAST + 1):
    for col in range(1, 6):
        c = ac.cell(row=r, column=col)
        c.fill = INPUT_FILL
        c.font = f(color="0000FF")
        c.border = BOX
    ac.cell(row=r, column=1).number_format = DATE_FMT
    ac.cell(row=r, column=3).number_format = "0"
    ac.cell(row=r, column=4).number_format = "0.0"
    last_date = f'_xlfn.MAXIFS({W},{W},"<="&A{r})'
    first_date = f'_xlfn.MINIFS({W},{W},">0")'
    formulas = [
        (6, f'=IF(A{r}="","",WEEKNUM(A{r},21))', "0"),
        (7, f'=IF(OR(C{r}="",N(D{r})=0),"",C{r}/D{r})', "0.0"),
        (8, f'=IF(A{r}="","",IFERROR(INDEX({WB},MATCH(IF({last_date}>0,{last_date},{first_date}),{W},0)),Settings!$B$6))', "0.0"),
        (9, f'=IF(B{r}="","",IFERROR(INDEX({MET_RANGE},MATCH(B{r},{ACT_RANGE},0)),4))', "0.0"),
        (10, f'=IF(OR(A{r}="",C{r}="",I{r}=""),"",I{r}*H{r}*C{r}/60)', "#,##0"),
    ]
    for col, formula, fmt in formulas:
        c = ac.cell(row=r, column=col, value=formula)
        c.font = f()
        c.number_format = fmt
        c.border = BOX
ac["A2"] = dt.date(2026, 10, 1)
ac["B2"] = "Running"
ac["C2"] = 40
ac["D2"] = 7.2
ac["E2"] = "Example row - overwrite or delete"
ac["F1"].comment = Comment("ISO week number (weeks start on Monday).", "Tracker")
ac["I1"].comment = Comment("MET of the activity, from Settings. 4.0 if the activity is not in the list.", "Tracker")
ac["J1"].comment = Comment("Estimate: MET x weight (kg) x duration (h).", "Tracker")
dv_act = DataValidation(type="list", formula1="=ActivityList", showErrorMessage=True,
                        error="Choose an activity from the list (edit the list in Settings).")
dv_dur = DataValidation(type="decimal", operator="between", formula1="1", formula2="1440",
                        showErrorMessage=True, error="Duration in minutes (1-1440).")
dv_dist = DataValidation(type="decimal", operator="between", formula1="0", formula2="1000",
                         showErrorMessage=True, error="Distance in km (0-1000).")
dv_date2 = DataValidation(type="date", operator="between", formula1="DATE(2000,1,1)",
                          formula2="DATE(2100,12,31)", showErrorMessage=True,
                          error="Enter a valid date.")
for dv, rng in ((dv_act, "B"), (dv_dur, "C"), (dv_dist, "D"), (dv_date2, "A")):
    ac.add_data_validation(dv)
    dv.add(f"{rng}2:{rng}{LAST}")

# ------------------------------------------------------------------ Dashboard
db = wb.create_sheet("Dashboard", 1)
db.sheet_view.showGridLines = False
for col, w in zip("ABCDEFGHIJKLMN", [26, 12, 12, 24, 12, 3, 3, 14, 14, 14, 14, 14, 14, 14]):
    db.column_dimensions[col].width = w
db["A1"] = "Activity & Weight Dashboard"
db["A1"].font = f(bold=True, size=16, color=NAVY)
db["A2"] = '="Updated: "&TEXT(TODAY(),"dd/mm/yyyy")'
db["A2"].font = f(italic=True, color="595959")

AA, AC_, AD, AJ, AB = ("Activities!$A$2:$A$1000", "Activities!$C$2:$C$1000",
                       "Activities!$D$2:$D$1000", "Activities!$J$2:$J$1000",
                       "Activities!$B$2:$B$1000")
week_start = "TODAY()-WEEKDAY(TODAY(),2)+1"
cards_left = ("Weight", [
    ("Current weight (kg)", f'=IFERROR(INDEX({WB},MATCH(MAX({W}),{W},0)),"-")', "0.0"),
    ("Starting weight (kg)", f'=IFERROR(INDEX({WB},MATCH(_xlfn.MINIFS({W},{W},">0"),{W},0)),"-")', "0.0"),
    ("Change since start (kg)", '=IFERROR(B5-B6,"-")', "+0.0;-0.0;0.0"),
    ("Target weight (kg)", "=Settings!$B$4", "0.0"),
    ("Still to go (kg)", '=IFERROR(B5-B8,"-")', "+0.0;-0.0;0.0"),
    ("BMI", '=IFERROR(B5/Settings!$B$3^2,"-")', "0.0"),
    ("Last weigh-in", f'=IF(COUNT({W})=0,"-",MAX({W}))', DATE_FMT),
])
cards_right = ("Activity", [
    ("Minutes this week", f'=SUMIFS({AC_},{AA},">="&{week_start},{AA},"<="&TODAY())', "0"),
    ("Weekly goal (min)", "=Settings!$B$5", "0"),
    ("% of weekly goal", '=IFERROR(E5/E6,0)', "0%"),
    ("Sessions - last 30 days", f'=COUNTIFS({AA},">"&TODAY()-30,{AA},"<="&TODAY())', "0"),
    ("Minutes - last 30 days", f'=SUMIFS({AC_},{AA},">"&TODAY()-30,{AA},"<="&TODAY())', "0"),
    ("Km - last 30 days", f'=SUMIFS({AD},{AA},">"&TODAY()-30,{AA},"<="&TODAY())', "0.0"),
    ("Kcal - last 30 days", f'=SUMIFS({AJ},{AA},">"&TODAY()-30,{AA},"<="&TODAY())', "#,##0"),
])
for col, (title, items) in ((1, cards_left), (4, cards_right)):
    t = db.cell(row=4, column=col, value=title)
    t.font = f(bold=True, color="FFFFFF", size=11)
    t.fill = HEAD_FILL
    db.cell(row=4, column=col + 1).fill = HEAD_FILL
    for i, (label, formula, fmt) in enumerate(items, start=5):
        a = db.cell(row=i, column=col, value=label)
        b = db.cell(row=i, column=col + 1, value=formula)
        a.font = f()
        b.font = f(bold=True, size=11)
        b.number_format = fmt
        b.alignment = Alignment(horizontal="right")
        for c in (a, b):
            c.fill = CARD_FILL
            c.border = BOX
db.conditional_formatting.add("E7", CellIsRule(operator="greaterThanOrEqual", formula=["1"],
                                               font=Font(name=FONT, bold=True, color="008000")))

# totals per activity
TOP = 14
header_cells = ["Activity", "Sessions", "Minutes", "Km", "Kcal (est.)"]
for i, h in enumerate(header_cells, start=1):
    c = db.cell(row=TOP, column=i, value=h)
    c.font = f(bold=True, color="FFFFFF")
    c.fill = HEAD_FILL
    c.border = BOX
    c.alignment = Alignment(horizontal="center")
for k in range(len(ACTIVITIES)):
    r = TOP + 1 + k
    s = ACT_FIRST + k
    vals = [
        (f"=Settings!$A${s}", "@"),
        (f'=IF(A{r}="",0,COUNTIFS({AB},A{r}))', "0"),
        (f'=IF(A{r}="",0,SUMIFS({AC_},{AB},A{r}))', "0"),
        (f'=IF(A{r}="",0,SUMIFS({AD},{AB},A{r}))', "0.0"),
        (f'=IF(A{r}="",0,SUMIFS({AJ},{AB},A{r}))', "#,##0"),
    ]
    for col, (v, fmt) in enumerate(vals, start=1):
        c = db.cell(row=r, column=col, value=v)
        c.font = f()
        c.border = BOX
        if fmt != "@":
            c.number_format = fmt
tot = TOP + 1 + len(ACTIVITIES)
db.cell(row=tot, column=1, value="Total").font = f(bold=True)
for col, fmt in ((2, "0"), (3, "0"), (4, "0.0"), (5, "#,##0")):
    L = db.cell(row=TOP, column=col).column_letter
    c = db.cell(row=tot, column=col, value=f"=SUM({L}{TOP+1}:{L}{tot-1})")
    c.font = f(bold=True)
    c.number_format = fmt
for col in range(1, 6):
    db.cell(row=tot, column=col).border = BOX
    db.cell(row=tot, column=col).fill = CARD_FILL
db.cell(row=TOP - 1, column=1, value="Totals per activity (all time)").font = f(bold=True, size=11, color=NAVY)
db["H4"] = "Buttons appear here after running the SetupButtons macro (see Instructions)."
db["H4"].font = f(italic=True, color="595959", size=9)

# charts
sc = ScatterChart()
sc.title = "Weight (kg)"
sc.style = 2
sc.height, sc.width = 8, 17
sc.x_axis.title = "Date"
sc.y_axis.title = "kg"
sc.x_axis.number_format = "dd/mm/yy"
sc.x_axis.majorTimeUnit = "days"
sc.legend = None
xs = Reference(wt, min_col=1, min_row=2, max_row=LAST)
ys = Reference(wt, min_col=2, min_row=2, max_row=LAST)
ser = Series(ys, xs, title="Weight")
ser.marker.symbol = "circle"
ser.marker.size = 5
ser.graphicalProperties.line.solidFill = "2E75B6"
ser.graphicalProperties.line.width = 22000
ser.smooth = False
sc.series.append(ser)
sc.y_axis.scaling.min = None
sc.x_axis.delete = False
sc.y_axis.delete = False
db.add_chart(sc, "H6")

bc = BarChart()
bc.type = "bar"
bc.title = "Minutes per activity"
bc.style = 2
bc.height, bc.width = 8, 17
bc.legend = None
bc.gapWidth = 40
data = Reference(db, min_col=3, min_row=TOP, max_row=TOP + len(ACTIVITIES))
cats = Reference(db, min_col=1, min_row=TOP + 1, max_row=TOP + len(ACTIVITIES))
bc.add_data(data, titles_from_data=True)
bc.set_categories(cats)
bc.series[0].graphicalProperties.solidFill = "2E75B6"
bc.x_axis.delete = False
bc.y_axis.delete = False
db.add_chart(bc, "H23")

wb.calculation.fullCalcOnLoad = True
wb.active = 1  # open on the Dashboard
wb.save(OUT)
print("saved", OUT)
